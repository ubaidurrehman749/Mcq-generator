import base64
import csv
import html as htmllib
import io
import json
import random
import re
from collections import Counter
from pathlib import Path

import pandas as pd
import streamlit as st

st.set_page_config(page_title="MCQ Generator · Study Sparks AI", page_icon="✦", layout="wide")

OWNER_NAME = "Ubaid ur Rehman"
LET = "ABCD"


# ---------- Logo ----------
def logo_tag():
    mimes = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
             ".webp": "image/webp", ".svg": "image/svg+xml"}
    for ext, mime in mimes.items():
        p = Path(__file__).parent / f"logo{ext}"
        if p.exists():
            b64 = base64.b64encode(p.read_bytes()).decode()
            return f'<img class="logo" src="data:{mime};base64,{b64}" alt="Study Sparks AI logo">'
    return '<span class="spark">✦</span>'


# ---------- Styling ----------
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Bricolage+Grotesque:wght@600;800&family=DM+Sans:wght@400;500;700&display=swap');
html, body, [class*="css"] { font-family: 'DM Sans', sans-serif; }
#MainMenu, footer, header { visibility: hidden; }
.block-container { padding-top: 1.2rem; max-width: 1100px; }
.topbar { display:flex; justify-content:space-between; align-items:center; flex-wrap:wrap; gap:12px;
          padding-bottom:14px; border-bottom:1px solid #FADFCB; }
.brand { display:flex; align-items:center; gap:12px; font-family:'Bricolage Grotesque',sans-serif;
         font-weight:800; font-size:1.45rem; color:#1A1A2E; }
.brand .logo { height:46px; width:46px; object-fit:contain; border-radius:10px; }
.brand .spark { color:#F2600C; font-size:1.6rem; }
.brand b { color:#F2600C; }
.owner { background:#1A1A2E; color:#FFF7F0; padding:8px 16px; border-radius:999px; font-weight:500; font-size:.95rem; }
.owner b { color:#FF9F43; }
.hero h1 { font-family:'Bricolage Grotesque',sans-serif; font-weight:800; font-size:2.4rem;
           line-height:1.1; color:#1A1A2E; margin:1.2rem 0 .3rem; letter-spacing:-.02em; }
.hero p { color:#5A5870; font-size:1.1rem; margin-bottom:.8rem; }

.qmeta { display:flex; justify-content:space-between; color:#5A5870; font-weight:700; font-size:.92rem; margin:6px 0; }
.qcard { background:linear-gradient(135deg,#1A1A2E,#2B2D55); color:#fff; border-left:8px solid #F2600C;
         border-radius:20px; padding:26px 28px; font-size:1.3rem; line-height:1.55; margin:6px 0 16px;
         box-shadow:0 22px 38px -24px rgba(26,26,46,.7); }
.qcard small { display:block; color:#FF9F43; font-weight:700; margin-bottom:8px; font-size:.82rem; }

.st-key-quizbox div[role="radiogroup"] { gap:8px; }
.st-key-quizbox div[role="radiogroup"] > label { background:#fff; border:1.5px solid #FADFCB; border-radius:14px;
         padding:12px 16px; transition:border-color .15s, background .15s; width:100%; }
.st-key-quizbox div[role="radiogroup"] > label:hover { border-color:#F2600C; background:#FFF1E6; }

.rev { background:#fff; border:1px solid #FADFCB; border-left:6px solid #16A34A; border-radius:12px;
       padding:10px 14px; margin:6px 0; color:#1A1A2E; }
.rev.bad { border-left-color:#DC2626; }
.rev small { color:#5A5870; }

div.stButton > button, div.stDownloadButton > button, div.stFormSubmitButton > button {
    border-radius:12px; font-weight:700; padding:.6rem 1rem; }
div.stButton > button[kind="primary"], div.stDownloadButton > button {
    background:#F2600C; color:#fff; border:0; }
div.stButton > button[kind="primary"]:hover, div.stDownloadButton > button:hover { background:#D24E05; color:#fff; }
[data-testid="stMetricValue"] { color:#F2600C; font-family:'Bricolage Grotesque',sans-serif; }
textarea { border-radius:14px !important; }
.stTabs [data-baseweb="tab"] { font-weight:700; }
</style>
""", unsafe_allow_html=True)

st.markdown(f"""
<div class="topbar">
  <div class="brand">{logo_tag()}<span>Study Sparks <b>AI</b></span></div>
  <div class="owner">Created by <b>{OWNER_NAME}</b></div>
</div>
<div class="hero">
  <h1>MCQ Generator</h1>
  <p>Turn notes, PDFs and slides into a quiz. Practise with instant feedback, take a timed-style exam, and export a printable paper.</p>
</div>
""", unsafe_allow_html=True)

# ---------- File reading ----------
FILE_TYPES = ["pdf", "docx", "pptx", "xlsx", "xlsm", "csv", "txt", "md", "html", "htm",
              "rtf", "json", "png", "jpg", "jpeg", "webp", "bmp", "tiff"]


@st.cache_data(show_spinner=False)
def extract_text(name: str, data: bytes) -> str:
    ext = Path(name).suffix.lower()

    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)

    if ext == ".docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" ".join(c.text for c in row.cells))
        return "\n".join(parts)

    if ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                    parts.append(shape.text_frame.text)
            if slide.has_notes_slide:
                parts.append(slide.notes_slide.notes_text_frame.text)
        return "\n".join(parts)

    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"{ws.title}.")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    parts.append(", ".join(cells) + ".")
        return "\n".join(parts)

    if ext in (".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff"):
        from PIL import Image
        import pytesseract
        return pytesseract.image_to_string(Image.open(io.BytesIO(data)))

    text = data.decode("utf-8", errors="ignore")

    if ext in (".html", ".htm"):
        text = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", text)
        text = re.sub(r"(?s)<[^>]+>", " ", text)
        return htmllib.unescape(text)

    if ext == ".rtf":
        text = re.sub(r"\\'[0-9a-f]{2}", " ", text)
        text = re.sub(r"\\[a-z]+-?\d* ?", " ", text)
        return re.sub(r"[{}]", " ", text)

    if ext == ".json":
        try:
            vals = []

            def walk(o):
                if isinstance(o, dict):
                    for v in o.values():
                        walk(v)
                elif isinstance(o, list):
                    for v in o:
                        walk(v)
                elif isinstance(o, str):
                    vals.append(o)
            walk(json.loads(text))
            return "\n".join(vals)
        except Exception:
            return text

    return text


# ---------- Question generation ----------
STOP = set("""a about above after again all also am an and any are as at be because been before being below
between both but by can could did do does doing down during each few for from further had has have having he her
here hers him his how i if in into is it its just me more most my no nor not now of off on once only or other our
out over own same she should so some such than that the their them then there these they this those through to too
under until up very was we were what when where which while who whom why will with would you your""".split())

MIN_WORDS = 60
DEF = re.compile(r"^(.{3,60}?)\s+(is|are|means|refers to|is defined as|is called)\s+(.{12,})$", re.I)


def words_of(text):
    return re.findall(r"[\w'’-]+", text)


def split_sentences(text):
    text = re.sub(r"\s+", " ", text).strip()
    parts = re.split(r"(?<=[.!?])\s+", text)
    return [p for p in parts if 6 <= len(words_of(p)) <= 45]


def local_questions(text, n, qtype, diff, seed):
    rnd = random.Random(seed)
    sents = split_sentences(text)
    low = [w.lower() for w in words_of(text)]
    freq = Counter(w for w in low if w not in STOP and len(w) > 3 and not w.isdigit())
    if len(freq) < 8 or not sents:
        return []
    top = max(freq.values())
    vocab = [w for w, _ in freq.most_common(150)]

    scored = []
    for i, s in enumerate(sents):
        ws = [w.lower() for w in words_of(s)]
        sc = sum(freq.get(w, 0) / top for w in ws) / (len(ws) ** 0.5)
        if DEF.match(s.rstrip(".!?")):
            sc *= 1.4
        scored.append((sc, i, s))
    ranked = [s for _, _, s in sorted(scored, reverse=True)]

    def style(word, like):
        if like.isupper() and len(like) > 1:
            return word.upper()
        return word.capitalize() if like[0].isupper() else word

    def pick(kw, sentence_low, k):
        pool = [w for w in vocab if w != kw.lower() and w not in sentence_low]
        if diff != "Easy":
            pool = sorted(pool, key=lambda w: (abs(len(w) - len(kw)), rnd.random()))[:14]
        if len(pool) < k:
            return None
        return [style(w, kw) for w in rnd.sample(pool, k)]

    def_terms = []
    for s in sents:
        m = DEF.match(s.rstrip(".!?"))
        if m and len(words_of(m.group(1))) <= 5:
            def_terms.append(m.group(1).strip())

    kinds = {"Multiple choice": ["mcq"], "True / False": ["tf"], "Mixed (both)": ["mcq", "tf"]}[qtype]
    out = []
    for s in ranked:
        if len(out) >= n:
            break
        kind = kinds[len(out) % len(kinds)]
        ws = words_of(s)
        slow = {w.lower() for w in ws}
        cands = [w for w in ws if w.lower() in freq and len(w) > 3]
        if not cands:
            continue
        kw = max(cands, key=lambda w: freq[w.lower()])

        if kind == "tf":
            if rnd.random() < 0.5:
                out.append({"q": f"True or False: {s}", "options": ["True", "False"], "answer": 0,
                            "explain": f"True. {s}", "type": "tf"})
            else:
                fake = pick(kw, slow, 1)
                if not fake:
                    continue
                stmt = re.sub(rf"\b{re.escape(kw)}\b", fake[0], s, count=1)
                out.append({"q": f"True or False: {stmt}", "options": ["True", "False"], "answer": 1,
                            "explain": f"False. The correct sentence is: {s}", "type": "tf"})
            continue

        m = DEF.match(s.rstrip(".!?"))
        if m and len(words_of(m.group(1))) <= 5:
            term = m.group(1).strip()
            others = [t for t in dict.fromkeys(def_terms) if t.lower() != term.lower()]
            ds = rnd.sample(others, 3) if len(others) >= 3 else pick(term.split()[0], slow, 3)
            if not ds:
                continue
            q = f"Which term is described as: “{m.group(3).strip()}”?"
            opts, ans = [term] + ds, term
        else:
            ds = pick(kw, slow, 3)
            if not ds:
                continue
            q = "Fill in the blank: " + re.sub(rf"\b{re.escape(kw)}\b", "_____", s, count=1)
            opts, ans = [kw] + ds, kw
        rnd.shuffle(opts)
        out.append({"q": q, "options": opts, "answer": opts.index(ans), "explain": f"Source: {s}", "type": "mcq"})
    return out


def ai_questions(text, n, qtype, diff):
    import anthropic
    client = anthropic.Anthropic(api_key=st.secrets["ANTHROPIC_API_KEY"])
    kind = {"Multiple choice": "multiple-choice questions with exactly 4 options",
            "True / False": 'true/false questions whose options are ["True","False"]',
            "Mixed (both)": 'a mix of multiple-choice questions (exactly 4 options) and true/false questions '
                            '(options ["True","False"])'}[qtype]
    prompt = (f"Create exactly {n} {kind} from the study material below for a student. Difficulty: {diff.lower()}. "
              f"Rules: questions must be clear and standalone; exactly one correct option; plausible distractors; "
              f"never use 'all of the above' or 'none of the above'; explanation is 1-2 sentences saying why the "
              f"answer is right; use the same language as the material. Return ONLY a JSON array like "
              f'[{{"q":"...","options":["...","...","...","..."],"answer":0,"explanation":"..."}}] where answer is '
              f"the 0-based index of the correct option.\n\n<text>\n{text[:60000]}\n</text>")
    msg = client.messages.create(model="claude-sonnet-5-5", max_tokens=6000,
                                 messages=[{"role": "user", "content": prompt}])
    m = re.search(r"\[.*\]", msg.content[0].text, re.S)
    data = json.loads(m.group(0))
    out = []
    for d in data:
        try:
            opts = [str(o).strip() for o in d["options"]]
            ans = int(d["answer"])
            if not (2 <= len(opts) <= 4) or not (0 <= ans < len(opts)) or not str(d["q"]).strip():
                continue
            is_tf = [o.lower() for o in opts] == ["true", "false"]
            if not is_tf:
                correct = opts[ans]
                random.shuffle(opts)
                ans = opts.index(correct)
            out.append({"q": str(d["q"]).strip(), "options": opts, "answer": ans,
                        "explain": str(d.get("explanation", "")).strip(), "type": "tf" if is_tf else "mcq"})
        except Exception:
            continue
    return out


def ai_available():
    try:
        return "ANTHROPIC_API_KEY" in st.secrets
    except Exception:
        return False


def make_docx(title, qs):
    import docx
    from docx.shared import Pt
    d = docx.Document()
    d.add_heading(title, 0)
    d.add_paragraph(f"Name: ______________________    Date: ______________    Marks: ____ / {len(qs)}")
    d.add_paragraph("Study Sparks AI")
    for i, q in enumerate(qs, 1):
        p = d.add_paragraph()
        r = p.add_run(f"Q{i}. {q['q']}")
        r.bold = True
        r.font.size = Pt(11)
        for j, o in enumerate(q["options"]):
            d.add_paragraph(f"{LET[j]}. {o}").paragraph_format.left_indent = Pt(24)
    d.add_page_break()
    d.add_heading("Answer key", 1)
    for i, q in enumerate(qs, 1):
        line = f"Q{i}. {LET[q['answer']]}. {q['options'][q['answer']]}"
        if q.get("explain"):
            line += f"  ({q['explain']})"
        d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


SAMPLE = ("Photosynthesis is the process by which green plants, algae and some bacteria convert light energy into "
          "chemical energy. Chloroplasts are the organelles where photosynthesis takes place, and they contain a "
          "green pigment called chlorophyll. Chlorophyll absorbs sunlight, especially in the red and blue "
          "wavelengths, and reflects green light. In the light-dependent reactions, sunlight splits water "
          "molecules, releasing oxygen and producing the energy carriers ATP and NADPH. The Calvin cycle uses ATP "
          "and NADPH to turn carbon dioxide into glucose. Glucose is used for growth, stored as starch, or broken "
          "down during cellular respiration to release energy. Light intensity, carbon dioxide concentration and "
          "temperature are limiting factors that affect the rate of photosynthesis. Respiration is the process "
          "that releases stored energy from glucose inside the mitochondria. Photosynthesis is important because "
          "it produces the oxygen we breathe and forms the base of nearly every food chain.")

# ---------- State ----------
ss = st.session_state
ss.setdefault("text", "")
ss.setdefault("up_key", 0)
ss.setdefault("questions", [])
ss.setdefault("order", [])
ss.setdefault("pos", 0)
ss.setdefault("answers", {})
ss.setdefault("checked", set())
ss.setdefault("finished", False)
ss.setdefault("round", 0)


def start_quiz(order=None):
    ss.order = list(order) if order is not None else list(range(len(ss.questions)))
    ss.pos = 0
    ss.answers = {k: v for k, v in ss.answers.items() if k not in ss.order} if order is not None else {}
    ss.checked = set() if order is None else {c for c in ss.checked if c not in ss.order}
    ss.finished = False
    ss.round += 1


def load_sample():
    ss.text = SAMPLE


def clear_text():
    ss.text = ""
    ss.up_key += 1


def esc(t):
    return htmllib.escape(str(t))


def cur_key():
    return f"r_{ss.round}_{ss.order[ss.pos]}"


def save_answer():
    v = ss.get(cur_key())
    if v is not None:
        ss.answers[ss.order[ss.pos]] = v


def do_check():
    save_answer()
    if ss.order[ss.pos] in ss.answers:
        ss.checked.add(ss.order[ss.pos])


def next_q():
    save_answer()
    ss.pos += 1
    if ss.pos >= len(ss.order):
        ss.pos = len(ss.order) - 1
        ss.finished = True


def prev_q():
    save_answer()
    ss.pos = max(0, ss.pos - 1)


def finish_now():
    save_answer()
    ss.finished = True


tab_create, tab_quiz, tab_edit = st.tabs(["① Create quiz", "② Take quiz", "③ Edit & export"])

# ======================= CREATE =======================
with tab_create:
    st.text_input("Quiz title", value="My Study Quiz", key="deck")
    st.text_area("Paste text", key="text", height=200,
                 placeholder="Paste your notes or article here…", label_visibility="collapsed")
    files = st.file_uploader("Or upload files (PDF, Word, PowerPoint, Excel, CSV, TXT, images and more)",
                             type=FILE_TYPES, accept_multiple_files=True, key=f"up_{ss.up_key}")

    extracted = []
    for f in files or []:
        try:
            with st.spinner(f"Reading {f.name}…"):
                t = extract_text(f.name, f.getvalue())
            if t.strip():
                extracted.append(t)
                st.caption(f"✅ {f.name}: {len(words_of(t)):,} words found")
            else:
                st.warning(f"{f.name}: no readable text found. For scanned PDFs, upload pages as images.")
        except Exception as e:
            st.warning(f"{f.name}: could not be read ({e})")

    full_text = "\n\n".join([ss.text] + extracted).strip()
    n_words = len(words_of(full_text))
    st.caption(f"Total: {n_words:,} words")

    b1, b2, _ = st.columns([1, 1, 3])
    b1.button("Try sample", on_click=load_sample, use_container_width=True)
    b2.button("Clear", on_click=clear_text, use_container_width=True)

    c1, c2, c3 = st.columns(3)
    n_q = c1.slider("Number of questions", 5, 30, 10)
    qtype = c2.selectbox("Question type", ["Multiple choice", "True / False", "Mixed (both)"])
    diff = c3.select_slider("Difficulty", ["Easy", "Medium", "Hard"], value="Medium")
    d1, d2 = st.columns(2)
    engines = ["Quick (offline)"] + (["AI (smarter questions)"] if ai_available() else [])
    engine = d1.radio("Engine", engines, horizontal=True)
    d2.radio("Quiz mode", ["Practice (instant feedback)", "Exam (results at the end)"], key="mode", horizontal=False)

    if st.button("Generate quiz", type="primary", use_container_width=True):
        if n_words < MIN_WORDS:
            st.error(f"Add more text or upload a file. At least {MIN_WORDS} words are needed for a good quiz.")
        else:
            try:
                with st.spinner("Creating your questions…"):
                    if engine.startswith("AI"):
                        qs = ai_questions(full_text, n_q, qtype, diff)
                    else:
                        qs = local_questions(full_text, n_q, qtype, diff, random.randint(1, 10 ** 6))
                if not qs:
                    raise ValueError("Could not create questions from this text. Try longer text with full sentences.")
                ss.questions = qs
                ss.answers = {}
                start_quiz()
                msg = f"{len(qs)} questions are ready. Open the **Take quiz** tab to start."
                if len(qs) < n_q:
                    msg += f" Only {len(qs)} of {n_q} could be made from this text."
                st.success(msg)
            except Exception as e:
                st.error(f"Something went wrong: {e}")

# ======================= TAKE QUIZ =======================
with tab_quiz:
    Q = ss.questions
    practice = ss.get("mode", "Practice").startswith("Practice")

    if not Q:
        st.info("No quiz yet. Create one in the **Create quiz** tab.")

    elif ss.finished:
        order = ss.order
        ok = [qi for qi in order if ss.answers.get(qi) == Q[qi]["answer"]]
        wrong = [qi for qi in order if qi not in ok]
        skipped = [qi for qi in order if qi not in ss.answers]
        pct = round(len(ok) / len(order) * 100) if order else 0

        st.markdown("### Quiz complete 🎉")
        m1, m2, m3, m4 = st.columns(4)
        m1.metric("Score", f"{pct}%")
        m2.metric("Correct", f"{len(ok)}/{len(order)}")
        m3.metric("Wrong", len(wrong) - len(skipped))
        m4.metric("Skipped", len(skipped))
        st.progress(pct / 100)
        if pct >= 80:
            st.success("Excellent work. You know this material well.")
        elif pct >= 50:
            st.warning("Good effort. Review the questions you missed and try again.")
        else:
            st.error("Needs more revision. Read your notes again, then retry the missed questions.")

        r1, r2 = st.columns(2)
        r1.button("Restart full quiz", on_click=start_quiz, use_container_width=True)
        r2.button("Retry missed questions", on_click=start_quiz, args=(wrong,), type="primary",
                  use_container_width=True, disabled=not wrong)

        st.markdown("#### Review")
        for n_, qi in enumerate(order, 1):
            q = Q[qi]
            got = ss.answers.get(qi)
            good = got == q["answer"]
            with st.expander(f"{'✅' if good else '❌'} Q{n_}. {q['q'][:80]}{'…' if len(q['q']) > 80 else ''}"):
                st.markdown(f"**{q['q']}**")
                yours = f"{LET[got]}. {q['options'][got]}" if got is not None else "Not answered"
                st.markdown(f'<div class="rev {"" if good else "bad"}"><small>Your answer</small><br>{esc(yours)}</div>'
                            f'<div class="rev"><small>Correct answer</small><br>'
                            f'{LET[q["answer"]]}. {esc(q["options"][q["answer"]])}</div>', unsafe_allow_html=True)
                if q.get("explain"):
                    st.info(q["explain"])

    else:
        order, pos = ss.order, ss.pos
        qi = order[pos]
        q = Q[qi]
        checked = qi in ss.checked

        st.progress(pos / len(order))
        st.markdown(f'<div class="qmeta"><span>Question {pos + 1} of {len(order)}</span>'
                    f'<span>{len(ss.answers)} answered · {"Practice" if practice else "Exam"} mode</span></div>',
                    unsafe_allow_html=True)
        st.markdown(f'<div class="qcard"><small>{"TRUE / FALSE" if q["type"] == "tf" else "CHOOSE THE BEST ANSWER"}</small>'
                    f'{esc(q["q"])}</div>', unsafe_allow_html=True)

        with st.container(key="quizbox"):
            st.radio("Choose an option", options=list(range(len(q["options"]))),
                     format_func=lambda i: f"{LET[i]}. {q['options'][i]}",
                     index=ss.answers.get(qi), key=cur_key(), disabled=checked and practice,
                     label_visibility="collapsed")

        if practice and checked:
            if ss.answers.get(qi) == q["answer"]:
                st.success("Correct!")
            else:
                st.error(f"Not quite. The correct answer is {LET[q['answer']]}. {q['options'][q['answer']]}")
            if q.get("explain"):
                st.info(q["explain"])

        last = pos == len(order) - 1
        n1, n2, n3 = st.columns(3)
        n1.button("◀ Previous", on_click=prev_q, use_container_width=True, disabled=pos == 0)
        if practice and not checked:
            n2.button("Check answer", on_click=do_check, type="primary", use_container_width=True)
        else:
            n2.button("Finish quiz" if last else "Next ▶", on_click=next_q, type="primary", use_container_width=True)
        n3.button("Finish now", on_click=finish_now, use_container_width=True)

# ======================= EDIT & EXPORT =======================
with tab_edit:
    Q = ss.questions
    if not Q:
        st.info("No quiz yet. Create one in the **Create quiz** tab.")
    else:
        st.caption("Edit any cell, add rows at the bottom, or delete rows. Leave an option empty if you need fewer "
                   "than four. Then select **Save changes**.")
        rows = []
        for q in Q:
            o = q["options"] + [""] * (4 - len(q["options"]))
            rows.append({"Question": q["q"], "A": o[0], "B": o[1], "C": o[2], "D": o[3],
                         "Answer": LET[q["answer"]], "Explanation": q.get("explain", "")})
        df = pd.DataFrame(rows)
        edited = st.data_editor(
            df, num_rows="dynamic", use_container_width=True, key=f"editor_{len(Q)}_{ss.round}",
            column_config={"Answer": st.column_config.SelectboxColumn("Answer", options=list(LET), required=True)})

        if st.button("Save changes", type="primary"):
            new, skipped = [], 0
            for _, r in edited.iterrows():
                qtext = str(r["Question"]).strip() if pd.notna(r["Question"]) else ""
                opts = [(L, str(r[L]).strip()) for L in LET if pd.notna(r[L]) and str(r[L]).strip()]
                letter = r["Answer"] if pd.notna(r["Answer"]) else None
                idx = next((i for i, (L, _) in enumerate(opts) if L == letter), None)
                if not qtext or len(opts) < 2 or idx is None:
                    skipped += 1
                    continue
                vals = [v for _, v in opts]
                tf = [v.lower() for v in vals] == ["true", "false"]
                expl = str(r["Explanation"]).strip() if pd.notna(r["Explanation"]) else ""
                new.append({"q": qtext, "options": vals, "answer": idx, "explain": expl, "type": "tf" if tf else "mcq"})
            ss.questions = new
            start_quiz()
            st.success(f"Saved {len(new)} questions." + (f" {skipped} incomplete rows were skipped." if skipped else ""))

        st.markdown("#### Export")
        title = ss.get("deck", "My Study Quiz")
        fname = re.sub(r"[^\w\-]+", "_", title).strip("_") or "quiz"

        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(["Question", "A", "B", "C", "D", "Answer", "Explanation"])
        for q in Q:
            o = q["options"] + [""] * (4 - len(q["options"]))
            w.writerow([q["q"], *o, LET[q["answer"]], q.get("explain", "")])

        txt = []
        for i, q in enumerate(Q, 1):
            txt.append(f"Q{i}. {q['q']}")
            txt += [f"   {LET[j]}. {o}" for j, o in enumerate(q["options"])]
            txt.append("")
        txt.append("ANSWER KEY")
        txt += [f"Q{i}. {LET[q['answer']]}" for i, q in enumerate(Q, 1)]

        e1, e2 = st.columns(2)
        e1.download_button("Download exam paper (Word)", make_docx(title, Q), file_name=f"{fname}.docx",
                           mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                           use_container_width=True)
        e2.download_button("Download text version (.txt)", "\n".join(txt), file_name=f"{fname}.txt",
                           use_container_width=True)
        e3, e4 = st.columns(2)
        e3.download_button("Download CSV", buf.getvalue(), file_name=f"{fname}.csv", mime="text/csv",
                           use_container_width=True)
        e4.download_button("Download JSON", json.dumps(Q, ensure_ascii=False, indent=2),
                           file_name=f"{fname}.json", mime="application/json", use_container_width=True)
        st.caption("The Word file has the questions first and the answer key on the last page. It is ready to print.")

st.markdown(f"<hr><center style='color:#5A5870'>© Study Sparks AI · Created by {OWNER_NAME}</center>",
            unsafe_allow_html=True)
